#!/usr/bin/env python3
"""Publish Part 5 from frozen, attributed source text using the house template.

The source DOCX is not modified. Editorial corrections are explicit and recorded
in the private JSON source, with paragraph numbers from the supplied document.
"""
from pathlib import Path
from html import escape
from io import BytesIO
import argparse
import hashlib
import json
import os
import re
import subprocess

from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ContentStream, DictionaryObject, NameObject
from reportlab import rl_config
from reportlab.pdfgen import canvas
import build_trenches_not_silos as house
from apply_publication_house_style import Publication, register_fonts, make_cover, ROMAN, draw_footer

ROOT = Path(__file__).resolve().parents[1]
SLUG = 'hold-collide-come-apart'
FILENAME = 'Monderman_Insight_Hold_Collide_Come_Apart_2026-10-01.pdf'
DATA = ROOT / 'pdf-src/hold-collide-come-apart.json'
CORRECTIONS = {
    26: ('at the lowest running cost it will ever have.', 'at a low coordination cost.'),
    27: ('Only the fourth kind, where the shared understanding is kept alive on purpose, moves anything across.', 'The fourth kind, where the shared understanding is kept alive on purpose, makes that exchange a regular part of the work.'),
    42: ('After it the price keeps rising.', 'After it the price can rise sharply.'),
    61: ('the company can no longer do anything that needs more than one division.', 'the company struggles to do work that needs more than one division.'),
    72: ('a bend that nobody chose has nobody to review it.', 'a bend that nobody chose may have no clear owner for its review.'),
    74: ('and it has that choice while it is still cheap.', 'and seeing it early may leave cheaper choices open.'),
}


def extract(source):
    doc = Document(source)
    assert len(doc.paragraphs) == 87, 'Unexpected source edition'
    assert doc.paragraphs[2].text == 'Hold, Collide, Come Apart'
    paragraphs = []
    changes = []
    for number, p in enumerate(doc.paragraphs, 1):
        if not 9 <= number <= 74 or not p.text:
            continue
        markup = ''
        for run in p.runs:
            text = run.text
            if number in CORRECTIONS:
                old, new = CORRECTIONS[number]
                text = text.replace(old, new)
            item = escape(text)
            if run.font.superscript and text.isdigit():
                item = f'<sup><a href="#reference-{text}" aria-label="Reference {text}">{text}</a></sup>'
            elif run.italic:
                item = '<em>' + item + '</em>'
            if run.bold:
                item = '<strong>' + item + '</strong>'
            markup += item
        final = p.text
        if number in CORRECTIONS:
            old, new = CORRECTIONS[number]
            assert old in p.text and escape(new) in markup
            final = final.replace(old, new)
            changes.append(dict(paragraph=number, before=old, after=new))
        paragraphs.append(dict(number=number, text=final, html=markup))
    refs = [p.text for p in doc.paragraphs[75:84]]
    old = refs[8]
    refs[8] = ('On Sony’s joint answer: Sony’s announcement of 4 May 2004, on the U.S. launch of Connect; Apple’s announcement of 28 April 2003, on the launch of the iTunes Music Store. Sony’s announcement of 30 August 2007 set out the phase-out of CONNECT music services in North America and Europe, with timing varying by region and no earlier than March 2008. Martin Belam, “Sony finally ready to disconnect music store,” 4 February 2008, currybet.net, reproduces Sony’s European customer notice specifying 31 March 2008 as the closing date.')
    changes.append(dict(paragraph=84, before=old, after=refs[8]))
    data = dict(title=doc.paragraphs[2].text, subtitle=doc.paragraphs[3].text,
                standfirst=doc.paragraphs[4].text, author='Jason Adamson', date='October 2026',
                category='GOVERNANCE AND PERFORMANCE', part=5, paragraphs=paragraphs,
                references=refs, biography=doc.paragraphs[85].text,
                copyright=doc.paragraphs[86].text, read_minutes=25,
                source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), corrections=changes)
    DATA.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    return data


def figures():
    return json.loads((ROOT/'pdf-src/hold-collide-figures.json').read_text())


def figure_html(fig, index):
    panels = []
    for j, panel in enumerate(fig['panels']):
        content = f'<h3>{escape(panel[0])}</h3>'
        if fig.get('diagrams'):
            content += house.diagram(j)
        content += ''.join(f'<p>{escape(s)}</p>' for s in panel[1:])
        panels.append(f'<div class="trench-panel">{content}</div>')
    angle = '''<svg class="hold-angle" viewBox="0 0 600 130" role="img" aria-label="The trench as dug departs from the line as drawn. The angle is exaggerated."><path d="M12 114H588" stroke="#9CC4C9" stroke-width="2" stroke-dasharray="6 5"/><path d="M12 114 588 18" stroke="#0C6E78" stroke-width="3"/><path d="M160 90V114M560 22V114" stroke="#0E3A44" stroke-width="2"/><text x="260" y="27" fill="#0E3A44">the trench as dug</text><text x="330" y="106" fill="#556268">the line as drawn</text></svg>''' if fig.get('angle') else ''
    wedge = '''<div class="hold-cost"><span>cost of dealing with it</span><svg viewBox="0 0 600 22" aria-hidden="true"><path d="M0 20 600 0V22H0Z" fill="#9CC4C9"/></svg></div>''' if fig.get('cost_wedge') else ''
    return f'''<figure class="trench-figure hold-figure hold-figure-{index}" aria-labelledby="figure-{index}-title">
<div class="trench-figure-box"><h2 id="figure-{index}-title">{escape(fig['title'])}</h2>
<p class="trench-figure-deck">{escape(fig['subtitle'])}</p>
{angle}
<div class="trench-panels trench-panels-{len(panels)}">{''.join(panels)}</div>
{wedge}
<p class="trench-figure-note">{escape(fig['note'])}</p></div>
<figcaption>{escape(fig['caption'])}</figcaption></figure>'''


def article_body(data):
    visual = {fig['after']: (i, fig) for i, fig in enumerate(figures(), 1)}
    captions = {15, 18, 42, 60, 69}
    parts = []
    for p in data['paragraphs']:
        if p['number'] in captions:
            continue  # Included exactly once in the figure's accessible caption.
        parts.append(f'<p data-source-paragraph="{p["number"]}">{p["html"]}</p>')
        if p['number'] in visual:
            i, fig = visual[p['number']]
            parts.append(figure_html(fig, i))
    return '\n'.join(parts)


def reference_html(data):
    def linked(text):
        result = escape(text)
        for text, url in [
            ('ibm.com/history/personal-computer', 'https://www.ibm.com/history/personal-computer'),
            ('currybet.net', 'https://www.currybet.net/cbet_blog/2008/02/sony-finally-ready-to-disconne.php'),
        ]:
            result = result.replace(text, f'<a href="{url}">{text}</a>')
        return result
    return '\n'.join(f'<li id="reference-{i}" data-source-reference="{i}">{linked(text)}</li>' for i, text in enumerate(data['references'], 1))


def build_html(data, pages):
    template = (ROOT/'trenches-not-silos.html').read_text()
    old = json.loads((ROOT/'pdf-src/trenches-not-silos.json').read_text())
    head = template[:template.index('<main class="article-main"')]
    for a,b in [(old['title'],data['title']), (old['subtitle'],data['subtitle']), (old['standfirst'],data['standfirst']),
                ('trenches-not-silos.html',SLUG+'.html'), ('Monderman_Insight_Trenches_Not_Silos_2026-09-28.pdf',FILENAME),
                ('September 2026','October 2026'), ('Part 4','Part 5'), ('10 pages · 20-minute read',f'{pages} pages · 25-minute read')]:
        head = head.replace(a,b)
    head = head.replace('https://www.monderman.com/assets/brand/monderman-social-card.png?v=20260907-wordmark2', f'https://www.monderman.com/assets/research/{SLUG}-social.png')
    head = head.replace('content="Monderman">', f'content="{data["title"]} | Governance and Performance · Part 5">')
    head = head.replace('</head>', '<link rel="stylesheet" href="hold-collide-publication.css?v=20261001.1">\n</head>')
    head = re.sub(r'<header\b[\s\S]*?</header>', (ROOT/'site-shell/header.html').read_text().strip(), head, count=1)
    body = f'''<main class="article-main" id="main-content">
<article class="article-body">{article_body(data)}</article>
<section class="article-references" aria-labelledby="references-heading"><h2 class="references-heading" id="references-heading">References</h2>
<ol class="reference-list">{reference_html(data)}</ol></section>
<section class="article-about" aria-labelledby="about-author-heading"><h2 class="about-heading" id="about-author-heading">About the author</h2>
<p>{escape(data['biography'])}</p><p>{escape(data['copyright'])}</p></section>
<section class="article-further" aria-label="Article links"><p class="article-kicker">Keep reading</p>
<p><a href="{FILENAME}" target="_blank" rel="noopener noreferrer">Download the PDF →</a></p>
<p><a href="trenches-not-silos.html">Read Part 4: Trenches, Not Silos →</a></p>
<p><a href="research.html#governance-performance-title">Read the Governance and Performance series →</a></p></section></main>
'''
    tail = template[template.index('<footer class="footer mond-footer"'):]
    tail = re.sub(r'<footer\b[\s\S]*?</footer>', (ROOT/'site-shell/footer.html').read_text().strip(), tail, count=1)
    (ROOT/(SLUG+'.html')).write_text(head + body + tail)


def build_pdf(data):
    scratch = ROOT/'tmp/pdfs/hold-collide'; scratch.mkdir(parents=True, exist_ok=True)
    # Reuse the approved Part 4 typography and closing-page layout, not its text.
    house.article_body = article_body
    house.reference_html = reference_html
    print_source = house.print_html(data).replace('</head>', f'<link rel="stylesheet" href="{(ROOT/"hold-collide-publication.css").as_uri()}"></head>')
    print_source = print_source.replace('Organizations deliver at the speed of their administrative reality.', 'Clearer Insight. Stronger Performance.')
    source = scratch/'body.html'; source.write_text(print_source)
    subprocess.run([os.environ.get('NODE', 'node'), str(ROOT/'scripts/render_publication_body.cjs'), str(source), str(scratch/'body.pdf')], check=True)
    register_fonts(); rl_config.canvas_basefontname = ROMAN
    pub = Publication(FILENAME, data['category']+' · PART 5', data['title'], data['subtitle'], data['standfirst'], data['date'], 0, ())
    writer = PdfWriter(); writer.append(PdfReader(BytesIO(make_cover(pub))))
    for page in PdfReader(scratch/'body.pdf').pages:
        stream = BytesIO(); c = canvas.Canvas(stream, pagesize=(612,792))
        draw_footer(c, data['date'], len(writer.pages)+1); c.save()
        page.merge_page(PdfReader(BytesIO(stream.getvalue())).pages[0]); writer.add_page(page)
    # Remove ReportLab's unused default font so every embedded face is NHG.
    for page in writer.pages:
        resources = DictionaryObject(page['/Resources'].get_object())
        fonts = DictionaryObject(resources['/Font'].get_object())
        defaults = {name for name,font in fonts.items() if font.get_object().get('/BaseFont') == '/Helvetica'}
        if defaults:
            roman = next(name for name,font in fonts.items() if any(s in str(font.get_object().get('/BaseFont')) for s in ('55Rg','NHG55')))
            content = ContentStream(page.get_contents(), writer)
            for args, op in content.operations:
                if op == b'Tf' and args[0] in defaults: args[0] = roman
            for name in defaults: del fonts[name]
            resources[NameObject('/Font')] = fonts; page[NameObject('/Resources')] = resources; page.replace_contents(content)
    writer.add_metadata({'/Title':data['title'], '/Author':data['author'], '/Subject':data['subtitle'], '/Keywords':'Governance and Performance, Part 5, organizational structure'})
    with (ROOT/FILENAME).open('wb') as out: writer.write(out)
    return len(writer.pages)


if __name__ == '__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--source',type=Path); parser.add_argument('--extract-only',action='store_true')
    args=parser.parse_args()
    data=extract(args.source) if args.source else json.loads(DATA.read_text())
    if not args.extract_only:
        pages=build_pdf(data); build_html(data,pages)
        from generate_publication_social_cards import Card, render
        render(Card(SLUG,'Insight','Governance and Performance · Part 5',data['title'],data['subtitle'],data['date'],('Hold, Collide,','Come Apart')))
        print(json.dumps(dict(file=FILENAME,pages=pages,paragraphs=len(data['paragraphs']),references=len(data['references']),figures=len(figures()))))
