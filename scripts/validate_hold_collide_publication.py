#!/usr/bin/env python3
"""Bounded content, typography, geometry and link checks for Part 5."""
from pathlib import Path
from collections import Counter
from html.parser import HTMLParser
import argparse
import hashlib
import json
import re
import logging
import pdfplumber
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
PDF='Monderman_Insight_Hold_Collide_Come_Apart_2026-10-01.pdf'
SLUG='hold-collide-come-apart.html'
logging.getLogger('pdfminer').setLevel(logging.ERROR)
def compact(s): return re.sub(r'\s','',s)
class Text(HTMLParser):
    def __init__(self): super().__init__(); self.text=[];self.links=[];self.ids=[];self.figures=0
    def handle_data(self,s): self.text.append(s)
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if tag=='a': self.links.append(a.get('href',''))
        if a.get('id'): self.ids.append(a['id'])
        if tag=='figure':self.figures+=1

def validate(source=None):
    data=json.loads((ROOT/'pdf-src/hold-collide-come-apart.json').read_text())
    figures=json.loads((ROOT/'pdf-src/hold-collide-figures.json').read_text())
    html=(ROOT/SLUG).read_text();text=Text();text.feed(html)
    visible=compact(' '.join(text.text)); checks=[]
    def check(v,label): assert v,label;checks.append(label)
    check(len(data['paragraphs'])==61 and len(data['references'])==9,'61 paragraphs including 5 captions, 9 references')
    check(len(figures)==text.figures==5,'All five figures preserved')
    check(len(text.ids)==len(set(text.ids)),'Unique HTML IDs')
    if source:
        from docx import Document
        check(hashlib.sha256(source.read_bytes()).hexdigest()==data['source_sha256'],'Source DOCX SHA matches')
        original=Document(source)
        edits={e['paragraph']:e for e in data['corrections']}
        for p in data['paragraphs']:
            s=original.paragraphs[p['number']-1].text
            if p['number'] in edits:
                e=edits[p['number']];s=s.replace(e['before'],e['after'])
            check(s==p['text'],f'P{p["number"]}: source differs only by recorded correction')
        for i,ref in enumerate(data['references'],76):
            check(ref==(edits[i]['after'] if i in edits else original.paragraphs[i-1].text),f'Reference P{i} source fidelity')
    with pdfplumber.open(ROOT/PDF) as pdf:
        pages=[' '.join(l['text'] for l in p.extract_text_lines() if l['top']<730) for p in pdf.pages]
        pdftext=compact(' '.join(pages))
        check(len(pages)==16,'16-page edition')
        for i,p in enumerate(pdf.pages[1:],2):
            check(all(c['x0']>=58 and c['x1']<=554 for c in p.chars),f'Page{i} within margins')
            check(all(c['top']>=55 and c['bottom']<763 for c in p.chars),f'Page{i} clear of page edges')
            check(str(i) in (p.crop((500,730,612,792)).extract_text() or ''),f'Page{i} correct folio')
        check(pages[14].startswith('REFERENCES'),'References on separate page')
    for label,surface in [('HTML',visible),('PDF',pdftext)]:
        positions=[surface.find(compact(p['text'])) for p in data['paragraphs']]
        check(all(p>=0 for p in positions) and positions==sorted(positions),label+' every paragraph/caption retained in order')
        for r in data['references']:check(compact(r) in surface,label+' complete reference')
        for k in ['title','subtitle','standfirst','biography','copyright']:check(compact(data[k]) in surface,label+' '+k)
        check('127.0.0.1' not in surface and 'localhost' not in surface,label+' no local QA URLs')
    for fig in figures:
        strings=[fig[k] for k in ['title','subtitle','note','caption']]+[s for panel in fig['panels'] for s in panel]
        check(all(compact(s) in visible for s in strings),'HTML figure wording '+fig['title'])
        words=Counter(re.findall(r'\w+',' '.join(pages)))
        check(all(words[w]>=n for w,n in Counter(re.findall(r'\w+',' '.join(strings))).items()),'PDF figure wording '+fig['title'])
    reader=PdfReader(ROOT/PDF);fonts=set()
    for page in reader.pages:
        check(tuple(float(x) for x in page.mediabox[2:])==(612,792),'US Letter')
        for f in page['/Resources']['/Font'].get_object().values():
            f=f.get_object();desc=f.get('/FontDescriptor')
            if desc is None and '/DescendantFonts' in f:desc=f['/DescendantFonts'][0].get_object().get('/FontDescriptor')
            name=str(f.get('/BaseFont') or (desc.get_object().get('/FontName') if desc else ''));fonts.add(name)
            check(any(k in name for k in ['NHaas','NeueHaas','NHG']),'House font '+name)
            embedded=f.get('/Subtype')=='/Type3' and bool(f.get('/CharProcs'))
            check(embedded or (desc and any(k in desc.get_object() for k in ['/FontFile','/FontFile2','/FontFile3'])),'Embedded font '+name)
    for link in text.links:
        if link.startswith('#'):check(link[1:] in text.ids,'Internal anchor '+link)
        elif not re.match(r'^[a-z]+:',link):check((ROOT/link.split('#')[0].split('?')[0]).exists(),'Local link '+link)
    check('16 pages · 25-minute read' in html,'Accurate article metadata')
    check('Governance and Performance · Part 5' in html,'Part 5 label')
    check('https://www.monderman.com/'+SLUG in html,'Canonical URL')
    check('assets/research/hold-collide-come-apart-social.png' in html,'Dedicated social preview')
    body=' '.join(p['text'] for p in data['paragraphs'])
    check(not re.search(r'McKinsey|Deloitte|Accenture|Bain|Boston Consulting|Booz Allen',body,re.I),'No consultancy competitor names in body')
    print(json.dumps(dict(passed=True,checks=len(checks),pages=16,figures=5,references=9,fonts=sorted(fonts)),indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--source',type=Path)
    validate(parser.parse_args().source)
